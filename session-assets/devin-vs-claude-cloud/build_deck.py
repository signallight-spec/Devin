from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

FONT = "Yu Gothic"
NAVY = RGBColor(0x1A, 0x23, 0x33)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
INDIGO_BG = RGBColor(0xEE, 0xEE, 0xFB)
AMBER = RGBColor(0xD9, 0x77, 0x06)
AMBER_BG = RGBColor(0xFD, 0xF3, 0xE3)
GRAY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xF7, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x05, 0x9A, 0x5B)
RED = RGBColor(0xDC, 0x26, 0x26)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def rect(slide, x, y, w, h, fill=None, line=None, shadow=False, round_=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    if round_:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph = list of (txt, size, bold, color)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for txt, size, bold, color in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = FONT
            rPr = r._r.get_or_add_rPr()
            ea = rPr.find(qn('a:ea'))
            if ea is None:
                ea = rPr.makeelement(qn('a:ea'), {})
                rPr.append(ea)
            ea.set('typeface', FONT)
    return tb


def set_eastasia(shape_or_tb):
    """Ensure East-Asian font follows Latin font for all runs."""
    for p in shape_or_tb.text_frame.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            ea = rPr.find(qn('a:ea'))
            if ea is None:
                ea = rPr.makeelement(qn('a:ea'), {})
                rPr.append(ea)
            ea.set('typeface', FONT)


def header(slide, kicker, title, sub=None, accent=INDIGO):
    rect(slide, 0, 0, SW, Inches(0.10), fill=accent)
    text(slide, Inches(0.55), Inches(0.32), Inches(12), Inches(0.3),
         [[(kicker, 11, True, accent)]])
    text(slide, Inches(0.55), Inches(0.58), Inches(12.3), Inches(0.6),
         [[(title, 26, True, NAVY)]])
    if sub:
        text(slide, Inches(0.55), Inches(1.12), Inches(12.3), Inches(0.4),
             [[(sub, 12, False, GRAY)]])


def page_footer(slide, n):
    text(slide, Inches(12.35), Inches(7.08), Inches(0.7), Inches(0.3),
         [[(str(n), 10, False, GRAY)]], align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, h, items, size=13, gap=6, color=NAVY, bullet="■  ", bcolor=INDIGO):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            head, body = it
            paras.append([(bullet, size - 1, True, bcolor),
                          (head, size, True, color),
                          (("  " + body) if body else "", size, False, GRAY)])
        else:
            paras.append([(bullet, size - 1, True, bcolor), (it, size, False, color)])
    text(slide, x, y, w, h, paras, space_after=gap, line_spacing=1.06)


# ============ S1 Title ============
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, Inches(4.9), Inches(0.55), Inches(0.09), fill=INDIGO)
text(s, Inches(0.55), Inches(1.6), Inches(12), Inches(0.5),
     [[("比較検討レポート  |  2026年9月", 13, False, RGBColor(0xA5, 0xB4, 0xFC))]])
text(s, Inches(0.55), Inches(2.15), Inches(12.3), Inches(1.8),
     [[("AIコーディングエージェント選定", 40, True, WHITE)],
      [("Claude Code Cloud vs Devin", 40, True, WHITE)]], space_after=10)
text(s, Inches(0.55), Inches(4.15), Inches(12), Inches(0.6),
     [[("結論: プラットフォームの完成度・統制・レビュー機能で Devin が優位", 17, True, RGBColor(0xC7, 0xD2, 0xFE))]])
text(s, Inches(0.55), Inches(6.55), Inches(12), Inches(0.4),
     [[("資料作成: Devin (Cognition)  |  情報ソース: 公式ドキュメント・発表記事", 10, False, RGBColor(0x8B, 0x93, 0xA7))]])

# ============ S2 Executive summary ============
s = prs.slides.add_slide(BLANK)
header(s, "EXECUTIVE SUMMARY", "結論：Devinを推奨します")
y = Inches(1.55)
cards = [
    ("Devin が優れる点", INDIGO, INDIGO_BG,
     ["レビュー〜修正〜CIまで自動で閉じる業務基盤",
      "Devin Review＋Auto-Fixの独自機能",
      "ACU統制・Secrets・統合など企業運用が成熟"]),
    ("Claude が優れる点", AMBER, AMBER_BG,
     ["最難関タスクでのモデル性能（Fable系）",
      "既契約者の追加コストが小さい",
      "ProjectsのマルチエージェントUX"]),
    ("判定", GREEN, RGBColor(0xE6, 0xF6, 0xEE),
     ["実務運用の総合力では Devin が優位",
      "特に「レビューと修正ループ」の差が決定的",
      "Claudeは個人利用の補助線として併用も可"]),
]
cw = Inches(4.0)
for i, (t, c, bg, items) in enumerate(cards):
    x = Inches(0.55) + i * (cw + Inches(0.22))
    rect(s, x, y, cw, Inches(4.6), fill=bg, round_=True)
    rect(s, x, y, cw, Inches(0.09), fill=c)
    text(s, x + Inches(0.25), y + Inches(0.28), cw - Inches(0.5), Inches(0.4),
         [[(t, 16, True, c)]])
    bullets(s, x + Inches(0.25), y + Inches(0.85), cw - Inches(0.5), Inches(3.5),
            items, size=12.5, gap=10, bcolor=c)
text(s, Inches(0.55), Inches(6.5), Inches(12.3), Inches(0.6),
     [[("本資料は、2026-09-24 GAの「Claude Code クラウドセッション」と Devin を機能・性能・コストで比較した評価です。", 12, False, GRAY)]])
page_footer(s, 2)

# ============ S3 Product overview ============
s = prs.slides.add_slide(BLANK)
header(s, "OVERVIEW", "両者の立ち位置")
colw = Inches(6.05)
x1, x2 = Inches(0.55), Inches(6.75)
rect(s, x1, Inches(1.55), colw, Inches(5.1), fill=LIGHT, round_=True)
rect(s, x1, Inches(1.55), colw, Inches(0.09), fill=AMBER)
text(s, x1 + Inches(0.3), Inches(1.85), colw - Inches(0.6), Inches(0.4),
     [[("Claude Code Cloud（2026-09-24 GA）", 15, True, AMBER)]])
bullets(s, x1 + Inches(0.3), Inches(2.4), colw - Inches(0.6), Inches(4.0), [
    ("Claude Codeのクラウド実行版。", "PCを閉じても継続、複数タスク並列・自動PR作成"),
    ("Projects（9/17追加）。", "コーディネーター＋スレッド群＋共有メモリ"),
    ("サンドボックス実行。", "隔離環境・許可済みリポジトリのみGit操作"),
    ("課金。", "サブスク内包＋従量クレジット。Pro $100/Max $250のトライアル提供中"),
    ("モデル。", "Claude系のみ（Sonnet/Opus/Fable等）"),
], size=12.5, gap=9, bcolor=AMBER)

rect(s, x2, Inches(1.55), colw, Inches(5.1), fill=LIGHT, round_=True)
rect(s, x2, Inches(1.55), colw, Inches(0.09), fill=INDIGO)
text(s, x2 + Inches(0.3), Inches(1.85), colw - Inches(0.6), Inches(0.4),
     [[("Devin（Cognition）", 15, True, INDIGO)]])
bullets(s, x2 + Inches(0.3), Inches(2.4), colw - Inches(0.6), Inches(4.0), [
    ("自律ソフトウェアエンジニアの業務基盤。", "コーディングだけでなく非ソフト作業も"),
    ("専用フルVM。", "環境スナップショット（blueprint）で再現・高速起動"),
    ("Devin Review。", "PR解析・バグ検出・セキュリティ・マージ操作を一体化"),
    ("統合。", "Slack/Teams・Linear/Jira・スケジュール・Oncall"),
    ("モデル選択。", "SWE-2 / GPT-5.6 / Fable 5.1 / Fusion（2モデル分担）"),
], size=12.5, gap=9, bcolor=INDIGO)
page_footer(s, 3)

# ============ S4 Comparison table ============
s = prs.slides.add_slide(BLANK)
header(s, "COMPARISON", "機能比較")
rows = [
    ("観点", "Claude Code Cloud", "Devin", True),
    ("実行環境", "管理サンドボックス", "専用フルVM＋環境スナップショット", False),
    ("モデル選択", "Claude系のみ", "SWE-2 / GPT-5.6 / Fable / Fusion", False),
    ("PRレビュー", "専用製品なし", "Devin Review（解析・検出・操作一体）", False),
    ("修正ループ", "手動起点", "Auto-Fix：指摘→修正→CIまで自律", False),
    ("チャット統合", "限定的", "Slack/Teams ネイティブセッション", False),
    ("自動化", "限定的", "スケジュール・イベント・Oncall", False),
    ("コスト統制", "月次クレジット上限", "ACU：ユーザー/組織/セッション別上限", False),
    ("対象範囲", "コーディング中心", "開発＋非ソフト作業全般", False),
]
tx, ty = Inches(0.55), Inches(1.55)
tw = Inches(12.25)
rh = Inches(0.56)
widths = [Inches(2.6), Inches(4.7), Inches(4.95)]
for ri, (c0, c1, c2, is_head) in enumerate(rows):
    y = ty + ri * rh
    bg = NAVY if is_head else (WHITE if ri % 2 else LIGHT)
    rect(s, tx, y, tw, rh, fill=bg)
    xs = [tx, tx + widths[0], tx + widths[0] + widths[1]]
    for xi, wv, txt in zip(xs, widths, (c0, c1, c2)):
        col = WHITE if is_head else NAVY
        if not is_head and xi == xs[0]:
            col = GRAY
        text(s, xi + Inches(0.18), y + Inches(0.09), wv - Inches(0.3), rh - Inches(0.15),
             [[(txt, 12.5 if is_head else 12, is_head or xi == xs[0], col)]],
             anchor=MSO_ANCHOR.MIDDLE)
page_footer(s, 4)

# ============ S5 Benchmark ============
s = prs.slides.add_slide(BLANK)
header(s, "PERFORMANCE", "ベンチマークで見る実力差")
s.shapes.add_picture("/home/ubuntu/deck/benchmark.png",
                     Inches(0.75), Inches(1.6), width=Inches(8.0))
rect(s, Inches(9.05), Inches(1.6), Inches(3.75), Inches(4.9), fill=INDIGO_BG, round_=True)
text(s, Inches(9.3), Inches(1.85), Inches(3.3), Inches(0.4),
     [[("読み取り", 14, True, INDIGO)]])
bullets(s, Inches(9.3), Inches(2.35), Inches(3.3), Inches(4.0), [
    ("SWE-2は主力指標で最上位勢に肉薄", "（Fable 5.1比64%安、GPT-6 Astra比約1/4コスト）"),
    ("弱点はTerminal-Bench 4", "（高度な自律タスク。Fable 55.8% vs SWE-2 27.3%）"),
    ("Devinはモデルを選べる", "ため、難タスクはFable等を選ぶ逃げ道がある"),
], size=11.5, gap=9)
text(s, Inches(0.75), Inches(6.7), Inches(12), Inches(0.5),
     [[("出典: Cognition「Introducing SWE-2」2026-09-10。各モデル最良effort設定の公表値。", 10.5, False, GRAY)]])
page_footer(s, 5)

# ============ S6 Devin advantages ============
s = prs.slides.add_slide(BLANK)
header(s, "WHY DEVIN", "Devinが優位な5つの理由")
adv = [
    ("レビュー基盤", "Devin Reviewは解析・バグ検出・セキュリティ・チャット・マージ操作まで一体化した独立製品。レビュー自体が自動化の対象になる"),
    ("自律修正ループ", "Auto-Fixで「レビュー指摘→修正→CI再実行」が人を介さず回り、マージ可能品質まで自律到達"),
    ("実行環境の完成度", "フルVM＋環境スナップショット。ブラウザ・Androidエミュレータ等の重い検証も可能"),
    ("統合と自動化", "Slack/Teamsネイティブ・Linear/Jira・スケジュール・イベント駆動・Oncall等、業務導線が揃う"),
    ("企業統制", "ACUメータリング・ユーザー別/組織別上限・Secrets・請求タグ・RBACと、管理機能が成熟"),
]
y = Inches(1.6)
for i, (t, d) in enumerate(adv):
    yy = y + i * Inches(1.02)
    rect(s, Inches(0.55), yy, Inches(0.42), Inches(0.42), fill=INDIGO, round_=True)
    text(s, Inches(0.55), yy + Inches(0.03), Inches(0.42), Inches(0.35),
         [[(str(i + 1), 15, True, WHITE)]], align=PP_ALIGN.CENTER)
    text(s, Inches(1.15), yy - Inches(0.02), Inches(3.0), Inches(0.4),
         [[(t, 15, True, INDIGO)]])
    text(s, Inches(4.15), yy + Inches(0.02), Inches(8.6), Inches(0.95),
         [[(d, 12, False, NAVY)]], line_spacing=1.05)
page_footer(s, 6)

# ============ S7 Fair view of Claude ============
s = prs.slides.add_slide(BLANK)
header(s, "FAIRNESS", "Claudeの強みと向き不向き")
x1, x2 = Inches(0.55), Inches(6.75)
rect(s, x1, Inches(1.55), colw, Inches(4.9), fill=AMBER_BG, round_=True)
text(s, x1 + Inches(0.3), Inches(1.8), colw - Inches(0.6), Inches(0.4),
     [[("Claudeが向くケース", 14, True, AMBER)]])
bullets(s, x1 + Inches(0.3), Inches(2.3), colw - Inches(0.6), Inches(4.0), [
    "既にPro/Max契約があり追加費を抑えたい個人",
    "Terminal-Bench 4級の難しい自律タスク（Fable系が強い）",
    "CLI→クラウドの同一ツール体験を求める場合",
    "Projectsのコーディネーター型並列作業が好み",
], size=12, gap=8, bcolor=AMBER)
rect(s, x2, Inches(1.55), colw, Inches(4.9), fill=LIGHT, round_=True)
text(s, x2 + Inches(0.3), Inches(1.8), colw - Inches(0.6), Inches(0.4),
     [[("評価上の注意", 14, True, GRAY)]])
bullets(s, x2 + Inches(0.3), Inches(2.3), colw - Inches(0.6), Inches(4.0), [
    "専用レビュー製品・Auto-Fix相当がなく、レビュー品質は人依存",
    "モデルがClaude系に限定（コスト/特性の選択肢なし）",
    "企業統制（予算上限・統合・監査）はDevinより浅い",
    "サンドボックスはVMほど自由度が高くない可能性",
], size=12, gap=8, bcolor=GRAY)
page_footer(s, 7)

# ============ S8 Conclusion ============
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, Inches(2.1), Inches(0.55), Inches(0.09), fill=INDIGO)
text(s, Inches(0.55), Inches(0.9), Inches(12), Inches(0.4),
     [[("CONCLUSION", 13, True, RGBColor(0xA5, 0xB4, 0xFC))]])
text(s, Inches(0.55), Inches(1.35), Inches(12.3), Inches(0.8),
     [[("Devin を推奨します", 36, True, WHITE)]])
text(s, Inches(0.55), Inches(2.45), Inches(12.3), Inches(2.2), [
    [("■  ", 13, True, INDIGO), ("「書く」だけでなく「レビュー・修正・CI・統制」まで含む業務基盤の総合力が上回る", 14, False, WHITE)],
    [("■  ", 13, True, INDIGO), ("モデル選択の自由度：安価なSWE-2で量を捌き、必要時は最上位モデルへ切替可", 14, False, WHITE)],
    [("■  ", 13, True, INDIGO), ("導入実績の当社向き機能：Slack連携・自動化・Secrets管理・コスト可視化", 14, False, WHITE)],
], space_after=14, line_spacing=1.15)
rect(s, Inches(0.55), Inches(5.15), Inches(12.25), Inches(1.35), fill=RGBColor(0x24, 0x30, 0x47), round_=True)
text(s, Inches(0.85), Inches(5.4), Inches(11.6), Inches(0.9), [
    [("併用案: ", 13, True, RGBColor(0xFC, 0xD3, 0x4D)),
     ("Devinを主軸に据えつつ、Claude契約者は個人作業の補助線・トライアルクレジットで難タスクの性能比較を行うのも合理的です。", 12.5, False, RGBColor(0xD1, 0xD5, 0xDB))],
], line_spacing=1.15)
text(s, Inches(0.55), Inches(6.85), Inches(12), Inches(0.4),
     [[("評価対象: Claude Code Cloud GA (2026-09-24発表) / Devin 現行機能", 10, False, RGBColor(0x8B, 0x93, 0xA7))]])

prs.save("/home/ubuntu/deck/Devin_vs_ClaudeCloud.pptx")
print("saved")
